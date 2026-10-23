"""
Tests for the book pipeline.

Fast tests (no --gpu needed):
  - parse_docx correctness
  - safe_name
  - --list-styles exit code

GPU tests (--gpu required):
  - End-to-end: DOCX → WAV + MP3
  - Resume: pre-existing section WAV is skipped
"""
import subprocess
import sys
from pathlib import Path

import pytest
import torch

from conftest import REPO_ROOT, make_test_docx
from tts_utils import save_wav_sf, load_wav   # portable WAV I/O (torchaudio.save unusable on ROCm)

FAKE_SR = 24000
_TOKEN_CAP = 1000
_S3_TOKEN_RATE = 25
TRUNC_SAMPLES = int(_TOKEN_CAP / _S3_TOKEN_RATE * FAKE_SR)  # exactly 40 s

SRC    = REPO_ROOT / "src"
PYTHON = sys.executable


# GPU profile tests were removed with `load_gpu_profile`; they checked profile parsing but did
# not validate measured render rates. Rate behavior is covered by `test_render_rate.py`.


# ── Post-scan safety net tests (no GPU needed) ───────────────────────────────

def test_post_scan_detects_and_replaces_truncated_chunk(tmp_path):
    """Post-scan must delete any truncated chunk and re-render it before concatenating.

    🔴 THIS TEST USED TO RE-IMPLEMENT THE POST-SCAN INSTEAD OF CALLING IT — twelve lines
    copied out of `finish_section` into the test body, `cpath(i).unlink()` included. So it
    proved the copy worked and said nothing whatsoever about the shipped code, and when that
    `unlink` turned out to be the one call in `book_to_audio` that 5.0-M6 never routed, this
    test was green through the whole life of the bug and went on being green while a real
    twelve-chapter render died on it (7.3.8). A test that reimplements its subject cannot
    fail with it.

    It drives `generate_section` now. Plaintext here, deliberately: the vault half is
    `test_m6_e2e.test_the_post_scan_can_DELETE_a_truncated_chunk_through_the_vault`, which is
    where the routing is what's under test. This one guards the behavior on a normal box.
    """
    import book_to_audio
    book_to_audio.torch = torch          # torch is imported lazily by main(); same as test_m6_e2e

    section = {"index": 0, "title": "Ch", "filename": "00_Ch",
               "text": "First sentence here. Second sentence here."}

    class _PoolThatTruncatesOnce:
        """The pool contract, no model: chunk 0 comes back 40 s long the first time."""
        sr = FAKE_SR

        def __init__(self):
            self.seen = []

        def run_chunks(self, items, params, save, total, on_event=None, chapter=None, cancel=None):
            for index, _text in items:
                n = TRUNC_SAMPLES if (index == 0 and index not in self.seen) else FAKE_SR
                self.seen.append(index)
                save(index, torch.zeros(1, n), self.sr)
            return 0, 0.0

        def submit_chunks(self, items, params, save, total, on_event=None, chapter=None, cancel=None):
            return self.run_chunks(items, params, save, total, on_event=on_event,
                                   chapter=chapter, cancel=cancel) if items else None

        def collect_chunks(self, batch, stall_timeout_s=None):
            return batch if batch is not None else (0, 0.0)

    import tts_utils
    out_dir = tmp_path / "out"
    out_dir.mkdir()                      # run_job's `vault_fs.mkdir(out_dir)`, done by hand
    pool = _PoolThatTruncatesOnce()
    out_path, dur, _nc, _nt = book_to_audio.generate_section(
        section, pool, tts_utils.GenParams(), out_dir, tmp_path / "chunks")

    assert pool.seen.count(0) == 2, (
        f"chunk 0 was rendered {pool.seen.count(0)} time(s) — the post-scan did not delete and "
        f"re-render the truncated chunk")
    chunks_dir = tmp_path / "chunks" / "00_Ch"
    banked = sorted(chunks_dir.glob("chunk_*.wav"))
    remaining = [p.name for p in banked if tts_utils.is_truncated_wav(p)]
    assert remaining == [], f"post-scan left a truncated chunk banked: {remaining}"
    # One second per chunk — the 40 s truncated one must not be in there.
    assert dur > 0 and load_wav(str(out_path))[0].shape[-1] == len(banked) * FAKE_SR, (
        "the truncated chunk reached the section WAV")


# ── MP3 duration tests (no GPU needed) ───────────────────────────────────────

def test_total_mp3_duration_sums_files(tmp_path, monkeypatch):
    """total_mp3_duration calls ffprobe per file and sums the results."""
    import subprocess
    mp3_dir = tmp_path / "mp3"
    mp3_dir.mkdir()
    (mp3_dir / "ch01.mp3").write_bytes(b"")
    (mp3_dir / "ch02.mp3").write_bytes(b"")

    durations = {"ch01.mp3": "1800.0\n", "ch02.mp3": "2700.5\n"}

    def mock_run(cmd, **kwargs):
        out = durations.get(Path(cmd[-1]).name, "")
        return subprocess.CompletedProcess(cmd, 0, stdout=out, stderr="")

    monkeypatch.setattr(subprocess, "run", mock_run)

    from book_to_audio import total_mp3_duration
    assert total_mp3_duration(mp3_dir) == pytest.approx(4500.5)


def test_total_mp3_duration_no_files(tmp_path):
    """Returns None when the mp3 directory is empty."""
    from book_to_audio import total_mp3_duration
    mp3_dir = tmp_path / "mp3"
    mp3_dir.mkdir()
    assert total_mp3_duration(mp3_dir) is None


@pytest.mark.skipif(not __import__("shutil").which("ffmpeg"), reason="ffmpeg required")
def test_wav_to_mp3_raises_on_bad_input(tmp_path):
    """A failed encode must raise, not return a path to a missing/partial MP3."""
    from tts_utils import wav_to_mp3
    bad = tmp_path / "not_audio.wav"
    bad.write_bytes(b"this is definitely not a wav file")
    with pytest.raises(RuntimeError):
        wav_to_mp3(bad)


def test_is_valid_wav(tmp_path):
    """Readable WAV → True; garbage (interrupted write) → False; missing → False."""
    from tts_utils import is_valid_wav
    good = tmp_path / "good.wav"
    save_wav_sf(str(good), torch.zeros(1, 2400), 24000)
    bad = tmp_path / "bad.wav"
    bad.write_bytes(b"RIFF garbage from a killed process")
    assert is_valid_wav(good) is True
    assert is_valid_wav(bad) is False
    assert is_valid_wav(tmp_path / "missing.wav") is False


# ── M2.5: valid_or_remove must not destroy on a failed READ ──────────────────
#
# Six callers, every one of them looping over every chunk of a render. The unlink is
# DELIBERATE for the case it was written for — a process killed mid-write leaves a
# truncated WAV, and treating it as "done" wedges the run forever. It is catastrophic for
# the case it could not tell apart from that one, because `is_valid_wav` asks soundfile,
# and libsndfile reports "I could not read this file" and "these bytes are not a WAV" as
# the SAME exception.

def test_valid_or_remove_still_deletes_a_genuinely_corrupt_chunk(tmp_path):
    """The behavior the docstring calls deliberate, and it stays. The bytes WERE read;
    they are simply not a WAV. That is a positive verdict, and it licenses the delete."""
    from tts_utils import valid_or_remove
    bad = tmp_path / "chunk_000.wav"
    bad.write_bytes(b"RIFF garbage from a killed process")
    assert valid_or_remove(bad) is False
    assert not bad.exists(), "a chunk we READ and found to be garbage must be regenerated"


def test_valid_or_remove_preserves_a_chunk_it_could_not_read(tmp_path, monkeypatch):
    """The M2.5 site. An OS-level read failure — EIO, ENOSPC's aftermath, a Windows sharing
    violation from the UI poller, and one day a wrong vault key — used to fall through
    `except Exception: return False` inside is_valid_wav and DELETE the chunk. Six callers
    loop over every chunk of the render, so one bad read took out the whole workspace in a
    single pass, behind a progress bar, one log line per file.

    The file must survive BYTE-FOR-BYTE. `unlink` is not the only way to destroy it: the
    callers answer False by re-rendering, which overwrites — so a bare `return False` would
    preserve the inode and lose the data just the same. That is why it raises."""
    import builtins
    import tts_utils

    good = tmp_path / "chunk_000.wav"
    save_wav_sf(str(good), torch.zeros(1, 2400), 24000)
    before = good.read_bytes()

    real_open = builtins.open

    def deny(path, *a, **kw):
        if Path(path) == good:
            raise OSError(5, "Input/output error")     # the disk, not the format
        return real_open(path, *a, **kw)

    monkeypatch.setattr(builtins, "open", deny)
    with pytest.raises(tts_utils.ChunkUnreadable):
        tts_utils.valid_or_remove(good)
    monkeypatch.undo()

    assert good.exists(), "a chunk we could not READ must never be unlinked"
    assert good.read_bytes() == before


def test_valid_or_remove_absent_is_not_corrupt(tmp_path):
    """ABSENT ≠ CORRUPT. A chunk that was never rendered has nothing to destroy: it is
    simply not banked, and saying so is not a claim about anything's contents."""
    from tts_utils import valid_or_remove
    assert valid_or_remove(tmp_path / "never_rendered.wav") is False


def test_print_book_duration_over_one_hour(tmp_path, monkeypatch, capsys):
    """Durations >= 1 hour are printed as Xh Ym."""
    import book_to_audio
    monkeypatch.setattr(book_to_audio, "total_mp3_duration", lambda p: 7380.0)  # 2h 3m
    book_to_audio.print_book_duration(tmp_path)
    assert "2h 3m" in capsys.readouterr().out


def test_print_book_duration_under_one_hour(tmp_path, monkeypatch, capsys):
    """Durations < 1 hour are printed as minutes only (no 'h')."""
    import book_to_audio
    monkeypatch.setattr(book_to_audio, "total_mp3_duration", lambda p: 2700.0)  # 45m
    book_to_audio.print_book_duration(tmp_path)
    out = capsys.readouterr().out
    assert "45m" in out
    assert "h" not in out


def test_print_book_duration_no_mp3s(tmp_path, monkeypatch, capsys):
    """Prints nothing when total_mp3_duration returns None."""
    import book_to_audio
    monkeypatch.setattr(book_to_audio, "total_mp3_duration", lambda p: None)
    book_to_audio.print_book_duration(tmp_path)
    assert capsys.readouterr().out == ""


# ── Unit tests (no GPU) ───────────────────────────────────────────────────────

def test_safe_name_basic():
    from book_to_audio import safe_name
    assert safe_name("Chapter One") == "Chapter_One"


def test_safe_name_strips_special_chars():
    from book_to_audio import safe_name
    assert "?" not in safe_name("What? Really!")


def test_safe_name_truncates():
    from book_to_audio import safe_name
    long_title = "A" * 100
    assert len(safe_name(long_title)) <= 50


def test_parse_docx_sections(tmp_path):
    """Minimal 2-chapter DOCX produces 2 sections with correct titles."""
    docx = tmp_path / "test.docx"
    make_test_docx(docx, num_chapters=2)

    from book_to_audio import parse_docx
    sections = parse_docx(docx)

    assert len(sections) == 2
    assert sections[0]["title"] == "Chapter 1"
    assert sections[1]["title"] == "Chapter 2"
    assert sections[0]["index"] == 0
    assert sections[1]["index"] == 1
    # Body text present in section text
    assert "chapter 1" in sections[0]["text"].lower()
    assert "chapter 2" in sections[1]["text"].lower()


def test_a_scene_break_typed_into_a_real_docx_never_reaches_the_model(tmp_path):
    """The whole path a book takes: a .docx with a "***" paragraph and a scene heading under it,
    through the real parser, to the chunks that are rendered."""
    from docx import Document
    import book_to_audio

    doc = Document()
    doc.add_heading("Chapter 1", level=1)
    for para in ("She closed the door behind her.", "***", "10 a.m. the next morning",
                 "Marcus was already at his desk."):
        doc.add_paragraph(para)
    path = tmp_path / "scene.docx"
    doc.save(str(path))

    (section,) = book_to_audio.parse_docx(path)
    assert "***" in section["text"], "the parser is expected to keep the book's text as written"
    assert book_to_audio.section_chunks(section) == [
        "Chapter 1. She closed the door behind her. 10 a.m. the next morning. "
        "Marcus was already at his desk."]


def test_parse_docx_filenames_are_safe(tmp_path):
    """Section filenames are safe for filesystem use."""
    docx = tmp_path / "test.docx"
    make_test_docx(docx, num_chapters=2)

    from book_to_audio import parse_docx
    sections = parse_docx(docx)
    for s in sections:
        assert "/" not in s["filename"]
        assert "\\" not in s["filename"]
        assert s["filename"].startswith(f"{s['index']:02d}_")


def test_parse_docx_toc_skipped(tmp_path):
    """A 'Table of Contents' heading and its body are skipped."""
    from docx import Document
    from book_to_audio import parse_docx

    doc = Document()
    doc.add_heading("Table of Contents", level=2)
    doc.add_paragraph("Chapter 1...........1")
    doc.add_paragraph("Chapter 2...........5")
    doc.add_heading("Chapter 1", level=2)
    doc.add_paragraph("Real body text here.")
    path = tmp_path / "toc.docx"
    doc.save(str(path))

    sections = parse_docx(path)
    assert len(sections) == 1
    assert sections[0]["title"] == "Chapter 1"
    # TOC entries must not appear in the section text
    assert "Chapter 1...........1" not in sections[0]["text"]


def _wrap_in_content_control(paragraphs, gallery=None, placeholder=False):
    """Move these paragraphs into one block content control, the way Word stores it: a plain one as
    in Doug's "Come Find Me", a Word-built table of contents when `gallery` is given, or one still
    showing its grey prompt when `placeholder` is true."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    sdt, pr, content = OxmlElement("w:sdt"), OxmlElement("w:sdtPr"), OxmlElement("w:sdtContent")
    if placeholder:
        pr.append(OxmlElement("w:showingPlcHdr"))
    if gallery:
        obj, gal = OxmlElement("w:docPartObj"), OxmlElement("w:docPartGallery")
        gal.set(qn("w:val"), gallery)
        obj.append(gal)
        pr.append(obj)
    else:
        ident = OxmlElement("w:id")
        ident.set(qn("w:val"), "-1984997786")
        pr.append(ident)
    sdt.append(pr)
    sdt.append(content)
    paragraphs[0]._p.addprevious(sdt)
    for p in paragraphs:
        content.append(p._p)


def test_parse_docx_reads_chapters_inside_a_content_control(tmp_path):
    """Doug, 2026-10-22: Chapters 12 to 25 of "Come Find Me" sat inside one hidden content control.
    The parser read only paragraphs directly in the body, so the render spoke "Chapter 12" with
    no body, skipped every chapter in the control, and went on at Part Three."""
    from docx import Document
    from book_to_audio import parse_docx

    doc = Document()
    doc.add_heading("Part Two", level=1)
    doc.add_heading("Chapter 12", level=2)
    inside = [doc.add_paragraph("Savannah had never been west of Alabama."),
              doc.add_heading("Chapter 13", level=2),
              doc.add_paragraph("The road ran on past Shreveport.")]
    doc.add_heading("Part Three", level=1)
    doc.add_paragraph("Back home again.")
    _wrap_in_content_control(inside)
    path = tmp_path / "control.docx"
    doc.save(str(path))

    sections = parse_docx(path)
    assert [s["title"] for s in sections] == ["Part Two", "Chapter 12", "Chapter 13", "Part Three"]
    assert "west of Alabama" in sections[1]["text"]
    assert "past Shreveport" in sections[2]["text"]
    assert "Back home again" in sections[3]["text"]


def test_parse_docx_does_not_narrate_an_unfilled_content_control(tmp_path):
    """A content control left from a template still shows Word's grey prompt. That prompt is not
    the book: it is skipped, whether the control holds whole paragraphs or sits inside one."""
    from docx import Document
    from docx.oxml import OxmlElement
    from book_to_audio import parse_docx

    doc = Document()
    doc.add_heading("Chapter 1", level=2)
    _wrap_in_content_control([doc.add_paragraph("Click or tap here to enter text.")], placeholder=True)
    para = doc.add_paragraph("She waited. ")
    run = para.add_run("Choose an item.")
    sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    pr = OxmlElement("w:sdtPr")
    pr.append(OxmlElement("w:showingPlcHdr"))
    sdt.append(pr)
    sdt.append(content)
    run._r.addprevious(sdt)
    content.append(run._r)
    path = tmp_path / "placeholder.docx"
    doc.save(str(path))

    (section,) = parse_docx(path)
    assert "She waited." in section["text"]
    assert "Click or tap" not in section["text"]
    assert "Choose an item" not in section["text"]

    # The flag is an on/off value: switched off, the control holds real text and is read.
    from docx.oxml.ns import qn
    pr.find(qn("w:showingPlcHdr")).set(qn("w:val"), "0")
    doc.save(str(path))
    (section,) = parse_docx(path)
    assert "Choose an item" in section["text"]


def test_parse_docx_reads_text_in_a_content_control_inside_a_paragraph(tmp_path):
    """"Come Find Me" Chapter 1's subtitle, "Presentable", was in a content control inside its
    paragraph. The parser took only the runs directly in the paragraph and lost it."""
    from docx import Document
    from docx.oxml import OxmlElement
    from book_to_audio import parse_docx

    doc = Document()
    doc.add_heading("Chapter 1", level=2)
    para = doc.add_paragraph()
    run = para.add_run("Presentable")
    sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    sdt.append(OxmlElement("w:sdtPr"))
    sdt.append(content)
    run._r.addprevious(sdt)
    content.append(run._r)
    doc.add_paragraph("Everyone in Mercer Falls agreed.")
    path = tmp_path / "inline.docx"
    doc.save(str(path))

    (section,) = parse_docx(path)
    assert "Presentable" in section["text"]


def test_parse_docx_reads_table_text_in_reading_order(tmp_path):
    """Text in a table cell is part of the book too (a letter or a sign set in a box)."""
    from docx import Document
    from book_to_audio import parse_docx

    doc = Document()
    doc.add_heading("Chapter 1", level=2)
    doc.add_paragraph("Before the note.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Meet me at the pier."
    doc.add_paragraph("After the note.")
    path = tmp_path / "table.docx"
    doc.save(str(path))

    (section,) = parse_docx(path)
    text = section["text"]
    assert text.index("Before the note") < text.index("Meet me at the pier") < text.index("After the note")


def test_parse_docx_still_skips_a_word_built_table_of_contents(tmp_path):
    """Word's own table of contents is a content control. It stays out whole, and it must not
    swallow a dedication between it and the first chapter."""
    from docx import Document
    from book_to_audio import parse_docx

    doc = Document()
    toc = [doc.add_paragraph("Contents"), doc.add_paragraph("Chapter 1 .......... 1")]
    doc.add_paragraph("For Grandma.")
    doc.add_heading("Chapter 1", level=2)
    doc.add_paragraph("Real body text here.")
    _wrap_in_content_control(toc, gallery="Table of Contents")
    path = tmp_path / "toc_control.docx"
    doc.save(str(path))

    sections = parse_docx(path)
    assert [s["title"] for s in sections] == ["Front Matter", "Chapter 1"]
    assert "For Grandma" in sections[0]["text"]
    assert ".........." not in "".join(s["text"] for s in sections)


def test_list_styles_exits_zero(tmp_path):
    """--list-styles should exit 0 and print style info without loading the model."""
    docx = tmp_path / "test.docx"
    make_test_docx(docx, num_chapters=1)

    result = subprocess.run(
        [PYTHON, str(SRC / "book_to_audio.py"), str(docx), "--list-styles"],
        cwd=str(SRC), capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60,
    )
    assert result.returncode == 0, result.stderr
    combined = (result.stdout + result.stderr).lower()   # loguru logs to stderr
    assert "heading 2" in combined


# ── Worker-pool wiring (2.0-M2, no GPU — fake pool) ──────────────────────────
#
# These drive generate_section through a fake pool that writes chunk files in
# REVERSE completion order, each chunk's amplitude encoding its index. Proves the
# section WAV is concatenated in (chunk) index order regardless of the order
# workers finish — the core M2 ordering guarantee — without loading a model.

_POOL_CHUNK_LEN = 2400   # 0.1 s at FAKE_SR


class _FakePool:
    """Minimal stand-in for WorkerPool.run_chunks used by generate_section."""
    sr = FAKE_SR

    def run_chunks(self, items, params, save, total, on_event=None, chapter=None, cancel=None):
        # Write in reverse order; amplitude 0.1*(index+1) marks each chunk.
        for index, _text in reversed(items):
            save(index, torch.full((1, _POOL_CHUNK_LEN), 0.1 * (index + 1)), self.sr)
        return 0, 0.0

    # 7.1 split the pool contract into submit + collect so a book can queue the next chapter
    # behind the current one (WorkerPool.submit_chunks). Both are expressed through
    # `run_chunks` here so the subclasses below — which override only that — keep working:
    # this fake renders eagerly at submit time and `collect_chunks` hands back the result.
    def submit_chunks(self, items, params, save, total, on_event=None, chapter=None, cancel=None):
        if not items:
            return None
        return self.run_chunks(items, params, save, total,
                               on_event=on_event, chapter=chapter, cancel=cancel)

    def collect_chunks(self, batch, stall_timeout_s=None):
        return batch if batch is not None else (0, 0.0)


def _make_section(n_chunks_min=4):
    from tts_utils import split_into_chunks
    text = " ".join(f"Sentence number {i} lives in this test section here." for i in range(40))
    chunks = split_into_chunks(text)
    assert len(chunks) >= n_chunks_min, f"need >= {n_chunks_min} chunks, got {len(chunks)}"
    return {"index": 0, "title": "Ch", "filename": "00_Ch", "text": text}, len(chunks)


def _wire_book_globals():
    """generate_section uses book_to_audio's torch module global (set in main)."""
    import book_to_audio
    book_to_audio.torch = torch
    return book_to_audio


def _mark_as_banked_for(book_to_audio, chunks_dir, section):
    """What an earlier, interrupted render of this same chapter text leaves beside its chunks."""
    import tts_utils
    (chunks_dir / book_to_audio.CHUNKS_MARK).write_text(
        tts_utils.chunks_fingerprint(book_to_audio.section_chunks(section)), encoding="utf-8")


class _TextRecordingPool(_FakePool):
    """Keeps the text of every chunk it is handed, and which indexes."""
    def __init__(self):
        self.texts, self.indices = {}, []

    def run_chunks(self, items, params, save, total, on_event=None, chapter=None, cancel=None):
        for index, text in items:
            self.texts[index] = text
            self.indices.append(index)
        return super().run_chunks(items, params, save, total)


# ── A book never narrates an asterisk (Doug, 2026-10-12) ─────────────────────

_SCENE_CHAPTER = ("Chapter One.\n\n"
                  "She closed the door behind her and let out a long breath.\n\n"
                  "***\n\n"
                  "10 a.m. the next morning\n\n"
                  "Marcus was *already* at his desk when she walked in.\n\n"
                  "* * *\n\n"
                  "57 seconds earlier\n\n"
                  "The phone had rung twice.")


def test_a_scene_break_in_a_book_is_a_pause_and_its_asterisks_are_never_spoken(tmp_path):
    """Every scene break in a book made Chatterbox babble: the "***" line went to the model
    inside a chunk. Nothing with an asterisk in it may be handed to the pool, the scene heading
    is its own sentence, and an *emphasized* word is still read."""
    import tts_utils
    book_to_audio = _wire_book_globals()
    section = {"index": 0, "title": "Chapter One", "filename": "00_Ch", "text": _SCENE_CHAPTER}
    out_dir = tmp_path / "out"; out_dir.mkdir()
    pool = _TextRecordingPool()

    book_to_audio.generate_section(section, pool, tts_utils.GenParams(), out_dir, tmp_path / "chunks")

    spoken = " ".join(pool.texts[i] for i in sorted(pool.texts))
    assert "*" not in spoken, spoken
    assert spoken == ("Chapter One. She closed the door behind her and let out a long breath. "
                      "10 a.m. the next morning. Marcus was already at his desk when she walked in. "
                      "57 seconds earlier. The phone had rung twice.")


def test_a_book_ends_a_sentence_only_beside_a_scene_break():
    """A book file made from a PDF has one paragraph per printed line. Away from a scene break
    its paragraphs are joined exactly as before; the full stop goes only on the paragraph next to
    the break (the scene heading, and the line the scene above ended on)."""
    book_to_audio = _wire_book_globals()
    wrapped = ("It was the best of times and it was\n\nthe worst of times and nobody\n\n"
               "could say which\n\n***\n\nLater that day\n\nthe two of them met again\n\nby the river.")
    section = {"index": 0, "title": "Ch", "filename": "00_Ch", "text": wrapped}
    assert book_to_audio.section_chunks(section) == [
        "It was the best of times and it was the worst of times and nobody could say which. "
        "Later that day. the two of them met again by the river."]


def test_the_books_chunk_list_is_the_cleaned_one_everywhere(tmp_path):
    """The render and the progress line both cut the chapter through `section_chunks`, so the
    count that is logged is the count that is rendered."""
    book_to_audio = _wire_book_globals()
    section = {"index": 0, "title": "Chapter One", "filename": "00_Ch", "text": _SCENE_CHAPTER}
    chunks = book_to_audio.section_chunks(section)
    assert chunks and not any("*" in c for c in chunks)
    src = Path(book_to_audio.__file__).read_text(encoding="utf-8")
    assert src.count("split_into_chunks(") == 1, (
        "book_to_audio cuts a chapter in more than one place; every cut must go through "
        "section_chunks(), or one of them hands the model the text as written")


def test_chunks_banked_before_the_text_was_cleaned_are_rendered_again(tmp_path):
    """A chapter half-rendered by an older version has chunk files and no mark. Its chunk 1 is
    for the old cut of the text, so it must not be stitched into the new one."""
    import tts_utils
    book_to_audio = _wire_book_globals()
    section, total = _make_section()
    out_dir = tmp_path / "out"; out_dir.mkdir()
    chunks_dir = tmp_path / "chunks" / section["filename"]
    chunks_dir.mkdir(parents=True)
    MARKER = 0.95
    save_wav_sf(str(chunks_dir / "chunk_0001.wav"), torch.full((1, _POOL_CHUNK_LEN), MARKER), FAKE_SR)
    pool = _TextRecordingPool()

    out_path, *_ = book_to_audio.generate_section(
        section, pool, tts_utils.GenParams(), out_dir, tmp_path / "chunks")

    assert sorted(pool.indices) == list(range(total)), "every chunk is rendered again"
    combined, _sr = load_wav(str(out_path))
    seg1 = combined[0, _POOL_CHUNK_LEN:2 * _POOL_CHUNK_LEN]
    assert abs(seg1.mean().item() - MARKER) > 0.1, "the old chunk 1 was stitched into the chapter"


def test_chunks_banked_for_an_edited_chapter_are_rendered_again(tmp_path):
    """Same number of chunks, different words: the count check cannot see it, the mark can."""
    import tts_utils
    book_to_audio = _wire_book_globals()
    section, total = _make_section()
    out_dir = tmp_path / "out"; out_dir.mkdir()
    chunks_dir = tmp_path / "chunks" / section["filename"]
    chunks_dir.mkdir(parents=True)
    before = dict(section, text=section["text"].replace("Sentence number 3 ", "Sentence number 9 "))
    assert len(book_to_audio.section_chunks(before)) == total
    _mark_as_banked_for(book_to_audio, chunks_dir, before)
    save_wav_sf(str(chunks_dir / "chunk_0001.wav"), torch.full((1, _POOL_CHUNK_LEN), 0.95), FAKE_SR)
    pool = _TextRecordingPool()

    book_to_audio.generate_section(section, pool, tts_utils.GenParams(), out_dir, tmp_path / "chunks")

    assert sorted(pool.indices) == list(range(total))
    assert (chunks_dir / book_to_audio.CHUNKS_MARK).read_text(encoding="utf-8") == \
        tts_utils.chunks_fingerprint(book_to_audio.section_chunks(section))


def test_generate_section_concats_chunks_in_index_order(tmp_path):
    import tts_utils
    book_to_audio = _wire_book_globals()
    section, total = _make_section()

    out_dir = tmp_path / "out"; out_dir.mkdir()
    chunks_base = tmp_path / "chunks"

    out_path, dur, _nc, _nt = book_to_audio.generate_section(
        section, _FakePool(), tts_utils.GenParams(), out_dir, chunks_base)

    assert out_path.exists()
    combined, sr = load_wav(str(out_path))
    assert sr == FAKE_SR
    assert combined.shape[-1] == total * _POOL_CHUNK_LEN
    # Segment k must carry amplitude 0.1*(k+1) → chunks stitched in index order,
    # even though the fake pool wrote them last-to-first.
    for k in range(total):
        seg = combined[0, k * _POOL_CHUNK_LEN:(k + 1) * _POOL_CHUNK_LEN]
        assert abs(seg.mean().item() - 0.1 * (k + 1)) < 0.02, f"segment {k} out of order"


def test_generate_section_resume_keeps_existing_chunk(tmp_path):
    """A pre-existing chunk file is not regenerated — the pool is only handed the
    missing indices — and its content survives into the final concat."""
    import tts_utils
    book_to_audio = _wire_book_globals()
    section, total = _make_section()

    out_dir = tmp_path / "out"; out_dir.mkdir()
    chunks_base = tmp_path / "chunks"
    section_chunks = chunks_base / section["filename"]
    section_chunks.mkdir(parents=True)
    _mark_as_banked_for(book_to_audio, section_chunks, section)

    # Pre-seed chunk index 1 with a distinctive amplitude the fake pool never uses.
    MARKER = 0.95
    save_wav_sf(str(section_chunks / "chunk_0001.wav"),
            torch.full((1, _POOL_CHUNK_LEN), MARKER), FAKE_SR)

    handed = {}
    class _RecordingPool(_FakePool):
        def run_chunks(self, items, params, save, total, on_event=None, chapter=None, cancel=None):
            handed["indices"] = sorted(i for i, _ in items)
            return super().run_chunks(items, params, save, total)

    out_path, *_ = book_to_audio.generate_section(
        section, _RecordingPool(), tts_utils.GenParams(), out_dir, chunks_base)

    assert 1 not in handed["indices"], "resume: existing chunk 1 must not be regenerated"
    combined, _sr = load_wav(str(out_path))
    seg1 = combined[0, 1 * _POOL_CHUNK_LEN:2 * _POOL_CHUNK_LEN]
    assert abs(seg1.mean().item() - MARKER) < 0.02, "resume: pre-existing chunk 1 was overwritten"


@pytest.mark.parametrize("extension", [".opus", ".mp3", ".wav"])
def test_start_section_skips_completed_chapter(tmp_path, extension):
    """A resumed web or CLI render must leave encoded chapters and chapter WAVs alone."""
    import book_to_audio
    import tts_utils

    section, _total = _make_section()
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    chunks_base = tmp_path / "chunks"
    stem = section["filename"]
    assert not book_to_audio.section_rendered(out_dir, stem)

    chapter = (out_dir / "mp3" / f"{stem}{extension}" if extension != ".wav"
               else out_dir / f"{stem}.wav")
    chapter.parent.mkdir(exist_ok=True)
    chapter.write_bytes(b"completed chapter")

    result = book_to_audio.start_section(
        section, _FakePool(), tts_utils.GenParams(), out_dir, chunks_base)

    assert result == {"done": (out_dir / f"{stem}.wav", None, 0, 0.0)}
    assert not chunks_base.exists(), "a completed chapter must not queue or recreate chunks"


def test_web_resume_records_opus_only_chapter_as_done(tmp_path, monkeypatch):
    """A resumed chapter whose WAV was already encoded must reappear in the sidecar."""
    import book_state
    from book_to_audio import books_root
    from tts_utils import GenParams
    from vocalis.runtime.jobs import Job
    from vocalis.web import books as web_books

    section, _total = _make_section()
    output_root = tmp_path / "output"
    out_dir = books_root(output_root) / "Book"
    encoded = out_dir / "mp3" / f"{section['filename']}.opus"
    encoded.parent.mkdir(parents=True)
    encoded.write_bytes(b"completed chapter")

    class _GpuGate:
        def acquire(self, timeout):
            return True

        def release(self):
            pass

    probed = []
    monkeypatch.setattr(web_books, "OUTPUT_DIR", output_root)
    monkeypatch.setattr(web_books, "CHUNKS_ROOT", tmp_path / "chunks")
    monkeypatch.setattr(web_books, "parse_book", lambda path: [section])
    monkeypatch.setattr(web_books, "GPUS", _GpuGate())
    monkeypatch.setattr(web_books.rt_pool, "resolve_devices", lambda spec: ["cuda:0"])
    monkeypatch.setattr(web_books.rt_pool, "worker_note", lambda workers, devices: "test worker")
    monkeypatch.setattr(web_books.rt_pool, "get_pool", lambda devices, workers: object())
    monkeypatch.setattr(web_books.rt_jobs, "vram_ticker", lambda job, stop: None)
    monkeypatch.setattr(web_books, "_gpu_name", lambda: "test gpu")
    monkeypatch.setattr(web_books.gpu, "cuda_count", lambda devices: 1)
    monkeypatch.setattr(web_books.render_rate, "get", lambda key: None)

    def duration(path):
        probed.append(path)
        return 12.5

    monkeypatch.setattr(book_state, "mp3_duration", duration)
    job = Job(id="resume-test", docx=str(tmp_path / "Book.docx"), voice="",
              params=GenParams(), gpus="cuda:0", workers=1, out_name="Book")
    web_books.run_job(job)

    assert job.status == "done", job.error
    assert probed == [encoded]
    sidecar = book_state.load_sidecar(out_dir)
    assert sidecar["chapters"][0]["status"] == "done"
    assert sidecar["chapters"][0]["duration_s"] == 12.5
    assert not (tmp_path / "chunks").exists(), "resume must not queue GPU chunks"


def test_generate_section_heals_corrupt_chunk_on_resume(tmp_path):
    """A corrupt chunk file (process killed mid-write) must be regenerated on
    resume, not counted as done — otherwise it wedges the run forever: it exists
    so it's never re-rendered, and the final concat crashes on it every time."""
    import tts_utils
    book_to_audio = _wire_book_globals()
    section, total = _make_section()

    out_dir = tmp_path / "out"; out_dir.mkdir()
    chunks_base = tmp_path / "chunks"
    section_chunks = chunks_base / section["filename"]
    section_chunks.mkdir(parents=True)
    _mark_as_banked_for(book_to_audio, section_chunks, section)
    # chunk 1: valid pre-existing; chunk 2: corrupt garbage
    save_wav_sf(str(section_chunks / "chunk_0001.wav"),
            torch.full((1, _POOL_CHUNK_LEN), 0.5), FAKE_SR)
    (section_chunks / "chunk_0002.wav").write_bytes(b"garbage from a killed process")

    handed = {}
    class _RecordingPool(_FakePool):
        def run_chunks(self, items, params, save, total, on_event=None, chapter=None, cancel=None):
            handed["indices"] = sorted(i for i, _ in items)
            return super().run_chunks(items, params, save, total)

    out_path, *_ = book_to_audio.generate_section(
        section, _RecordingPool(), tts_utils.GenParams(), out_dir, chunks_base)

    assert 2 in handed["indices"], "corrupt chunk 2 must be regenerated"
    assert 1 not in handed["indices"], "valid chunk 1 must still resume"
    assert out_path.exists()                       # concat completed cleanly


# ── Integration tests (GPU required) ─────────────────────────────────────────

def test_generate_one_chapter(gpu_enabled, tmp_path, voice):
    """End-to-end: minimal 1-chapter DOCX produces a WAV and MP3."""
    docx    = tmp_path / "book.docx"
    out_dir = tmp_path / "out"
    make_test_docx(docx, num_chapters=1)

    result = subprocess.run(
        [PYTHON, str(SRC / "book_to_audio.py"),
         str(docx),
         "--voice",   voice,
         "--out-dir", str(out_dir)],
        cwd=str(SRC), capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=600,
    )
    assert result.returncode == 0, result.stderr

    # 3.x: the WAV master is transient — deleted once the MP3 lands.
    assert not list(out_dir.glob("*.wav")), "chapter WAV should be deleted after MP3 encode"

    # 🔴 GLOB THE DIRECTORY, NOT THE EXTENSION. Chapters are OPUS now (`book_to_audio.py:827`
    # calls `wav_to_opus`); the folder keeps the name `mp3/` so URLs already issued still
    # resolve, and the log line still says "Encoded MP3: 00_Chapter_1.opus". This assertion
    # was written when the two agreed, and when they stopped it went red on any box with a
    # GPU and stayed silent everywhere else — `gpu_enabled` skips it without hardware, which
    # is why a stale test survived. Matching `*` keeps it honest through the next change of
    # codec, which is not what this test is about.
    chapters = [p for p in (out_dir / "mp3").glob("*") if p.is_file()]
    assert len(chapters) == 1, f"Expected 1 encoded chapter, got {[m.name for m in chapters]}"
    assert chapters[0].suffix == ".opus", f"expected opus, got {chapters[0].name}"
    assert chapters[0].stat().st_size > 500


def test_book_artifacts_are_mp3_only(gpu_enabled, tmp_path, voice):
    """3.x: a finished book keeps ONLY the chapter MP3s — the WAV master and the
    chunk workspace are transient, deleted per-section as each MP3 lands (the old
    --delete-chunks flag is gone; cleanup is the behavior now)."""
    from book_to_audio import safe_name, CHUNKS_ROOT
    docx    = tmp_path / "mp3_only.docx"
    out_dir = tmp_path / "out"
    make_test_docx(docx, num_chapters=1)
    chunks_base = CHUNKS_ROOT / safe_name(docx.stem)

    try:
        result = subprocess.run(
            [PYTHON, str(SRC / "book_to_audio.py"),
             str(docx), "--voice", voice, "--out-dir", str(out_dir)],
            cwd=str(SRC), capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=600,
        )
        assert result.returncode == 0, result.stderr
        chapters = [p for p in (out_dir / "mp3").glob("*") if p.is_file()]
        assert len(chapters) == 1 and chapters[0].stat().st_size > 500
        assert not list(out_dir.glob("*.wav")), "chapter WAV must be deleted once its MP3 lands"
        assert not list(chunks_base.rglob("chunk_*.wav")) if chunks_base.exists() else True, \
            "chunk workspace must be cleaned per-section"
    finally:
        if chunks_base.exists():
            import shutil; shutil.rmtree(chunks_base)


def test_resume_skips_completed_section(gpu_enabled, tmp_path, voice):
    """
    Pre-populate section 0's WAV with a silent file.
    Running book_to_audio.py should skip section 0 and generate section 1 only.
    """
    import torch

    docx    = tmp_path / "resume.docx"
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    make_test_docx(docx, num_chapters=2)

    # Determine the section 0 filename that parse_docx will produce
    from book_to_audio import parse_docx, safe_name
    sections = parse_docx(docx)
    sec0_wav = out_dir / (sections[0]["filename"] + ".wav")
    # Located by STEM, not by extension — chapters are opus now (see the note in
    # test_generate_one_chapter). The directory keeps the name `mp3/`.
    sec1_stem = sections[1]["filename"]

    # Pre-create a valid silent WAV for section 0
    save_wav_sf(str(sec0_wav), torch.zeros(1, 2400), 24000)
    mtime_before = sec0_wav.stat().st_mtime

    result = subprocess.run(
        [PYTHON, str(SRC / "book_to_audio.py"),
         str(docx),
         "--voice",   voice,
         "--out-dir", str(out_dir)],
        cwd=str(SRC), capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=600,
    )
    assert result.returncode == 0, result.stderr

    # Section 0 must be untouched (a legacy WAV still counts as done)
    assert sec0_wav.stat().st_mtime == mtime_before, "Section 0 was regenerated — resume broken"

    # Section 1 must now exist — as its encoded chapter (3.x: the WAV master is transient)
    made = [q for q in (out_dir / "mp3").glob(sec1_stem + ".*") if q.is_file()]
    assert made, f"Section 1 was not generated (mp3/ holds {[q.name for q in (out_dir / 'mp3').glob('*')]})"
    assert made[0].stat().st_size > 500


# ── M3: no plaintext seams (5.0 §8) ──────────────────────────────────────────

def test_opening_a_book_writes_nothing_to_the_os_temp_dir(tmp_path, monkeypatch):
    """THE biggest leak the sweep found, and the easiest to miss because it looks like
    housekeeping. `open_docx` did mkstemp + copy2 — a byte-for-byte plaintext copy of THE
    ENTIRE BOOK into the OS temp directory, on EVERY parse. Every book click, every render,
    every regen. Outside the install root, outside cache_root(), outside every path 5.0
    encrypts, and outside anything anyone would think to wipe.

    It bought exactly one thing — reading a .docx while Word holds it open — and Word allows
    READS, so reading the bytes ourselves buys the same thing for nothing. python-docx takes
    a file-like, and the old code's own comment already conceded it "loads everything into
    memory" anyway.

    Asserted by watching tempfile itself, not by inspecting the source: a future edit that
    reintroduces a temp copy by another route still fails this."""
    import tempfile
    import book_to_audio

    docx = tmp_path / "My Private Novel.docx"
    make_test_docx(docx)

    calls = []
    for fn in ("mkstemp", "NamedTemporaryFile", "TemporaryDirectory", "mkdtemp"):
        real = getattr(tempfile, fn)
        monkeypatch.setattr(tempfile, fn,
                            lambda *a, _f=fn, _r=real, **kw: (calls.append(_f), _r(*a, **kw))[1])

    sections = book_to_audio.parse_docx(docx)

    assert sections, "the book must still actually parse"
    assert calls == [], f"the book was copied into the OS temp dir via tempfile.{calls}"


def test_a_book_can_be_read_while_something_else_holds_it_open(tmp_path):
    """The ONE thing the temp copy was buying. Word takes an exclusive WRITE lock but still
    permits reads — so read_book_bytes must work with a handle already open on the file,
    which on Windows is where a naive rewrite would fall over."""
    import book_to_audio

    docx = tmp_path / "locked.docx"
    make_test_docx(docx)
    with open(docx, "rb"):                       # someone else has it open
        sections = book_to_audio.parse_docx(docx)
    assert sections
